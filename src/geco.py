"""GECO DATA is authoritative; aggregate observed durations without skip imputation."""
import csv
import math
import string
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook
from src.data import CORPUS

IDENTITY=('WORD','PART','TRIAL','WORD_ID_WITHIN_TRIAL')
DVS={'gaze':'WORD_GAZE_DURATION','total':'WORD_TOTAL_READING_TIME'}


def workbook_rows(path):
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        it=book['DATA'].iter_rows(values_only=True)
        header=next(it)
        for values in it:
            yield dict(zip(header,values))
    finally:
        book.close()


def aggregate(observations):
    regions={}; seen=set(); participants=set(); diag=Counter()
    totals=defaultdict(lambda:Counter())
    numeric_regions=set(); boolean_regions=set()
    for r in observations:
        if isinstance(r['WORD'],bool):
            boolean_regions.add(str(r['WORD_ID']))
            r=dict(r,WORD='TRUE' if r['WORD'] else 'FALSE')
        if isinstance(r['WORD'],(int,float)) and not isinstance(r['WORD'],bool):
            numeric_regions.add(str(r['WORD_ID']))
            r=dict(r,WORD=format(r['WORD'],'g'))
        wid=str(r['WORD_ID']); pp=str(r['PP_NR'])
        ident=tuple(r[c] for c in IDENTITY)
        if wid in regions and regions[wid]!=ident:
            raise ValueError(f'Inconsistent WORD/position for {wid}: {ident} vs {regions[wid]}')
        if (pp,wid) in seen:
            raise ValueError(f'Duplicate participant/word: {pp}, {wid}')
        if not isinstance(r['WORD'],str) or not r['WORD'] or not r['WORD'].strip():
            raise ValueError(f'Empty/invalid reading region {wid}: {r["WORD"]!r}')
        regions[wid]=ident; seen.add((pp,wid)); participants.add(pp)
        t=totals[wid]; t['participants']+=1
        skip=r['WORD_SKIP']
        if skip not in (0,1):
            raise ValueError(f'Invalid skip status {wid}: {skip}')
        t['skip']+=int(skip); diag['skipped_observations']+=int(skip)
        for name,col in DVS.items():
            value=r[col]
            if value in ('.',None,''):
                diag[f'{name}_missing']+=1
            else:
                value=float(value)
                if not math.isfinite(value) or value<=0:
                    diag[f'{name}_nonpositive_or_invalid']+=1
                    continue
                t[f'{name}_n']+=1; t[f'{name}_sum']+=value
                diag[f'{name}_valid']+=1
                if skip:
                    diag[f'{name}_valid_and_skip']+=1
        diag['participant_rows']+=1
        if diag['participant_rows']%100000==0:
            print(f'GECO: read {diag["participant_rows"]:,} participant rows',flush=True)
    ordered=sorted(regions,key=lambda wid:tuple(int(x) for x in regions[wid][1:]))
    contexts=sorted({(int(regions[w][1]),int(regions[w][2])) for w in regions})
    ids={c:i+1 for i,c in enumerate(contexts)}
    rows=[]; keys=set()
    for wid in ordered:
        word,part,trial,zone=regions[wid]; part,trial,zone=map(int,(part,trial,zone))
        if (part,trial,zone) in keys:
            raise ValueError('Different WORD_IDs share a trial position')
        keys.add((part,trial,zone)); t=totals[wid]
        rows.append(dict(WORD_ID=wid,PART=part,TRIAL=trial,WORD_ID_WITHIN_TRIAL=zone,
            item=ids[part,trial],zone=zone,word=word,n_participants=t['participants'],
            n_gaze=t['gaze_n'],proportion_valid_gaze=t['gaze_n']/t['participants'],
            n_total=t['total_n'],skip_proportion=t['skip']/t['participants'],
            mean_gaze=t['gaze_sum']/t['gaze_n'] if t['gaze_n'] else '',
            mean_total=t['total_sum']/t['total_n'] if t['total_n'] else ''))
    gaps=[]
    for item in ids.values():
        context=[r for r in rows if r['item']==item]
        positions=[r['WORD_ID_WITHIN_TRIAL'] for r in context]
        if positions!=list(range(1,len(positions)+1)):
            gaps.append(dict(item=item,PART=context[0]['PART'],TRIAL=context[0]['TRIAL'],
                first_position=min(positions),last_position=max(positions),observed_regions=len(positions),
                missing_positions=sorted(set(range(min(positions),max(positions)+1))-set(positions))))
        # Internal zone is ordinal in the authoritative observed sequence. Preserve original IDs separately.
        for ordinal,r in enumerate(context,1): r['zone']=ordinal
    return rows,dict(diag,unique_regions=len(rows),trials=len(contexts),participants=len(participants),
                     invalid_regions=0,noncontiguous_trial_positions=gaps,numeric_excel_word_regions=len(numeric_regions),boolean_excel_word_regions=len(boolean_regions),scalar_conversion_ids=sorted(numeric_regions|boolean_regions),multiword_regions=sum(any(c.isspace() for c in r['word']) for r in rows),skip_proportion=diag['skipped_observations']/len(seen))


def frequency_lookup(path=CORPUS/'freqs/freqs-1.tsv'):
    counts=defaultdict(set)
    with open(path) as f:
        for code,order,token,count,_ in csv.reader(f,delimiter='\t'):
            if code.endswith('.word') and token and ' ' not in token:
                counts[token].add(int(count))
    return {w:next(iter(c)) for w,c in counts.items() if len(c)==1}


def add_lexical(rows,lookup):
    output=[]
    for r in rows:
        form=r['word'].strip(string.punctuation+'“”‘’…')
        count=lookup.get(form)
        output.append(dict(r,word_length=len(r['word']),frequency_form=form,
            frequency_count=count if count is not None else '',
            log_frequency=math.log1p(count) if count is not None else ''))
    return lag_columns(output,['word_length','log_frequency'])


def lag_columns(rows,columns):
    lookup={(int(r['item']),int(r['zone'])):r for r in rows}
    if len(lookup)!=len(rows):
        raise ValueError('Duplicate region keys')
    return [dict(r,**{f'prev_{c}':lookup.get((int(r['item']),int(r['zone'])-1),{}).get(c,'')
                       for c in columns}) for r in rows]


def material_diagnostic(rows,path=Path('data/geco/EnglishMaterial.xlsx')):
    """Optional comparison only: never used to filter or reconstruct DATA stimuli."""
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        values=book['ALL'].iter_rows(min_row=2,max_col=4,values_only=True)
        material={str(r[0]):r[3] for r in values if r[0] is not None}
    finally:
        book.close()
    data={r['WORD_ID']:r['word'] for r in rows}
    shared=set(data)&set(material)
    return dict(data_unique_ids=len(data),material_unique_ids=len(material),
        data_ids_absent_from_material=len(set(data)-set(material)),
        material_ids_absent_from_data=len(set(material)-set(data)),
        raw_word_disagreements=sum(data[k]!=material[k] for k in shared),
        policy='Diagnostic only; no inner join or stimulus modification')
